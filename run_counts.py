"""
Full pipeline: detect the 1 mm2 counting square in each photo, self-check it,
count beads inside it, and write the results table.

Per repeat the 5 photos are the 4 corner squares (4x4 cells) then the centre
square (5x5), centre last -> that ordering sets the self-check's expected grid.

Two folder layouts are supported. Old: a condition folder with r1/ and r2/
subfolders. New: all photos loose in the condition folder, taken in capture
order and split first-5 = r1, remainder = r2. If a folder doesn't hold exactly
10 photos (user took one more/less) it is still counted and a warning printed;
the overall average pools all squares so the odd photo can't distort beads/mL.

Outputs (results/):
  <condition>/unmasked/<repeat>/<name>.jpg            raw original photo
  <condition>/masked/<repeat>/<name>_masked.jpg      only the square visible
  <condition>/annotated/<repeat>/<name>_counted.jpg  masked + counted beads
  results_per_image.csv    count, ignored, bead area, detect + self-check info
  results_summary.csv      R1x5, R2x5, averages, beads/mL, total beads, recovery %
  and a printed summary table. All counts are per 1 mm2 square.
"""

import os
import glob
import csv
import shutil
import cv2
from count_beads import count_image, apply_square_mask
from detect_square import best_square, verify_subdivision

REPEATS = ["r1", "r2"]

VOL_PER_SQUARE_ML = 1e-4   # one large square (one photo) = 1e-4 mL

# Per-condition sample volume in mL, matched by a substring of the condition
# folder name (case-insensitive). Only list volumes that are fixed run-to-run;
# anything not listed is left BLANK in the Excel sheet for you to type in (the
# recovery formula then updates live). e.g. the transferred volume changes each
# experiment, so it is intentionally not listed here.
CONDITION_VOLUMES = {"bulk": 200.0}
# Condition whose total bead count is the 100% reference for recovery %
# (substring match, case-insensitive). Set to None to omit the recovery column.
REFERENCE_CONDITION = "bulk"


def volume_for(condition):
    lc = condition.lower()
    for key, vol in CONDITION_VOLUMES.items():
        if key.lower() in lc:
            return vol
    return None   # blank -> editable in the xlsx


HDR = ["condition", "R1_1", "R1_2", "R1_3", "R1_4", "R1_5",
       "R2_1", "R2_2", "R2_3", "R2_4", "R2_5",
       "R1_avg", "R2_avg", "overall_avg", "beads_per_mL",
       "sample_volume_mL", "total_beads", "recovery_pct"]


def write_summary_xlsx(summary_rows, hdr, path):
    """Write the summary as a real .xlsx where `sample_volume_mL` is an editable
    (yellow) cell and `total_beads` / `recovery_pct` are live Excel formulas, so
    changing a volume recalculates recovery automatically."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)

    col = {h: i + 1 for i, h in enumerate(hdr)}
    L = {h: get_column_letter(col[h]) for h in hdr}
    yellow = PatternFill("solid", fgColor="FFF2CC")

    # reference row (for recovery %); data rows start at spreadsheet row 2
    ref_row = None
    if REFERENCE_CONDITION:
        for i, row in enumerate(summary_rows, start=2):
            if REFERENCE_CONDITION.lower() in row["condition"].lower():
                ref_row = i
                break

    for i, row in enumerate(summary_rows, start=2):
        for h in hdr:
            cell = ws.cell(row=i, column=col[h])
            if h == "sample_volume_mL":
                cell.value = row[h] if row[h] != "" else None
                cell.fill = yellow            # highlight: type the volume here
            elif h == "total_beads":
                cell.value = f"={L['beads_per_mL']}{i}*{L['sample_volume_mL']}{i}"
            elif h == "recovery_pct":
                cell.value = (f"={L['total_beads']}{i}/{L['total_beads']}{ref_row}*100"
                              if ref_row else "")
                cell.number_format = "0.0"
            else:
                cell.value = row.get(h, "")
    for c in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    wb.save(path)

# Root folder that holds the condition folders. Chosen at run time via a folder
# picker (or passed as a command-line argument). OUT is set once BASE is known.
BASE = None
OUT = None


def pick_base_dir():
    """Return the conditions root folder: use a command-line argument if given,
    otherwise pop up a folder-selection dialog."""
    import sys
    if len(sys.argv) > 1:
        return sys.argv[1]
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    folder = filedialog.askdirectory(
        title="Select the folder that contains your condition folders")
    root.destroy()
    return folder


def find_conditions():
    """A condition folder is one that either has r1/r2 subfolders (old layout)
    or directly contains photos (new layout: all photos loose in the folder)."""
    conds = []
    for d in sorted(os.listdir(BASE)):
        full = os.path.join(BASE, d)
        if not os.path.isdir(full) or d == "results":
            continue
        has_repeats = any(os.path.isdir(os.path.join(full, r)) for r in REPEATS)
        if has_repeats or ordered_images(full):
            conds.append(d)
    return conds


def _is_source_photo(path):
    name = os.path.basename(path)
    return ("_annotated" not in name and "_counted" not in name
            and "_masked" not in name)


def ordered_images(folder):
    """All source photos directly in `folder`, in capture order (file mtime,
    tie-broken by filename). Our own output images are excluded."""
    paths = []
    for ext in ("*.jpg", "*.jpeg", "*.png"):
        paths += glob.glob(os.path.join(folder, ext))
    paths = [p for p in dict.fromkeys(paths) if _is_source_photo(p)]
    return sorted(paths, key=lambda p: (os.path.getmtime(p),
                                        os.path.basename(p)))


def images_in(condition, repeat):
    folder = os.path.join(BASE, condition, repeat)
    return sorted(p for p in glob.glob(os.path.join(folder, "*.jpg"))
                  if _is_source_photo(p))


def repeats_for(condition):
    """Return {"r1": [...paths...], "r2": [...]} for a condition, supporting
    both layouts. Old layout: read the r1/ and r2/ subfolders. New layout: take
    all loose photos in capture order and split first-5 = r1, remainder = r2.
    Extra/missing photos (not exactly 10) are tolerated: they still get counted,
    only a warning is printed so the odd folder can be eyeballed."""
    full = os.path.join(BASE, condition)
    if any(os.path.isdir(os.path.join(full, r)) for r in REPEATS):
        return {r: images_in(condition, r) for r in REPEATS}
    photos = ordered_images(full)
    r1, r2 = photos[:5], photos[5:]
    if len(photos) != 10:
        print(f"  WARNING: {condition} has {len(photos)} photos "
              f"(expected 10) -> r1={len(r1)}, r2={len(r2)}")
    return {"r1": r1, "r2": r2}


def collect_annotated(out_dir):
    """Gather every annotated image into results/all_annotated/ under a flat,
    sortable name  <condition>__<repeat>__pNN__<original>.jpg  so they can all be
    scrolled in a single folder, grouped by condition, then repeat, then photo
    order. Re-runnable: the folder is rebuilt from scratch each time."""
    dest = os.path.join(out_dir, "all_annotated")
    os.makedirs(dest, exist_ok=True)
    for stale in glob.glob(os.path.join(dest, "*.jpg")):
        os.remove(stale)                       # drop old files so nothing lingers
    n = 0
    for cond in sorted(os.listdir(out_dir)):
        cdir = os.path.join(out_dir, cond, "annotated")
        if not os.path.isdir(cdir):
            continue                           # skip all_annotated/ and non-conditions
        for rep in sorted(os.listdir(cdir)):
            rdir = os.path.join(cdir, rep)
            if not os.path.isdir(rdir):
                continue
            # annotated names are <original>_counted.jpg and the QS_#### numbers
            # run in capture order, so sorting by name gives the photo order.
            for idx, p in enumerate(sorted(glob.glob(os.path.join(rdir, "*.jpg"))),
                                    start=1):
                orig = os.path.basename(p).replace("_counted", "")
                shutil.copy2(p, os.path.join(
                    dest, f"{cond}__{rep}__p{idx:02d}__{orig}"))
                n += 1
    return dest, n


def main():
    global BASE, OUT
    BASE = pick_base_dir()
    if not BASE or not os.path.isdir(BASE):
        print("No folder selected - nothing to do.")
        return
    OUT = os.path.join(BASE, "results")
    print(f"Processing conditions in: {BASE}\n")

    per_image_rows = []
    summary_rows = []

    conditions = find_conditions()
    if not conditions:
        print("No condition folders (containing r1/r2) found in that folder.")
        return
    for cond in conditions:
        cond_counts = {}
        rep_paths = repeats_for(cond)
        for rep in REPEATS:
            counts = []
            paths = rep_paths[rep]
            # results/<condition>/{unmasked,masked,annotated}/<repeat>/
            dirs = {}
            for kind in ("unmasked", "masked", "annotated"):
                d = os.path.join(OUT, cond, kind, rep)
                os.makedirs(d, exist_ok=True)
                dirs[kind] = d

            for idx, path in enumerate(paths):
                img = cv2.imread(path)
                gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
                det = best_square(gray)
                # last photo of the 5 = centre square (5x5), others corners (4x4)
                n_cells = 5 if idx == len(paths) - 1 else 4
                ok, frac = verify_subdivision(gray, det, n_cells)
                r = count_image(path, region=det["corners"])

                name = os.path.splitext(os.path.basename(path))[0]
                # unmasked: raw original; masked: only the square visible;
                # annotated: masked image with counted beads drawn on.
                cv2.imwrite(os.path.join(dirs["unmasked"], name + ".jpg"), img)
                cv2.imwrite(os.path.join(dirs["masked"], name + "_masked.jpg"),
                            apply_square_mask(img, det["corners"]))
                cv2.imwrite(os.path.join(dirs["annotated"], name + "_counted.jpg"),
                            r["annotated"])
                counts.append(r["count"])
                per_image_rows.append(dict(
                    condition=cond, repeat=rep, image=os.path.basename(path),
                    square=("centre_5x5" if n_cells == 5 else "corner_4x4"),
                    count_per_mm2=r["count"], raw_maxima=r.get("raw_maxima", ""),
                    singles=r["singles"],
                    agg_beads=r["agg_beads"], agg_clumps=r["agg_clumps"],
                    ignored_aggregates=r["ignored"], bead_area_px=r["bead_area"],
                    detect_score=round(det["score"], 2),
                    selfcheck_ok=ok, selfcheck_frac=round(frac, 2)))
                flag = "" if ok else "  <-- SELF-CHECK FAILED, review"
                print(f"  {cond}/{rep}/{os.path.basename(path)}: "
                      f"{r['count']}/mm2 (ignored {r['ignored']}){flag}")
            cond_counts[rep] = counts

        r1, r2 = cond_counts.get("r1", []), cond_counts.get("r2", [])
        row = {"condition": cond}
        for i in range(5):
            row[f"R1_{i+1}"] = r1[i] if i < len(r1) else ""
        for i in range(5):
            row[f"R2_{i+1}"] = r2[i] if i < len(r2) else ""
        row["R1_avg"] = round(sum(r1) / len(r1), 1) if r1 else ""
        row["R2_avg"] = round(sum(r2) / len(r2), 1) if r2 else ""
        allc = r1 + r2
        mean_sq = sum(allc) / len(allc) if allc else 0
        row["overall_avg"] = round(mean_sq, 1)
        # concentration: one square = VOL_PER_SQUARE_ML, so beads/mL = mean/vol
        beads_per_ml = mean_sq / VOL_PER_SQUARE_ML
        vol = volume_for(cond)
        row["beads_per_mL"] = round(beads_per_ml)
        row["sample_volume_mL"] = vol if vol is not None else ""
        row["total_beads"] = round(beads_per_ml * vol) if vol is not None else ""
        summary_rows.append(row)

    # recovery %: each condition's total beads vs the reference condition's total
    # (static values for the CSV; the xlsx uses live formulas instead)
    ref_total = None
    if REFERENCE_CONDITION:
        for row in summary_rows:
            if (REFERENCE_CONDITION.lower() in row["condition"].lower()
                    and row["total_beads"] != ""):
                ref_total = row["total_beads"]
                break
    for row in summary_rows:
        row["recovery_pct"] = (round(100.0 * row["total_beads"] / ref_total, 1)
                               if ref_total and row["total_beads"] != "" else "")

    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "results_per_image.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(per_image_rows[0].keys()))
        w.writeheader(); w.writerows(per_image_rows)
    with open(os.path.join(OUT, "results_summary.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(summary_rows[0].keys()))
        w.writeheader(); w.writerows(summary_rows)
    write_summary_xlsx(summary_rows, HDR, os.path.join(OUT, "results_summary.xlsx"))
    dest, n_collected = collect_annotated(OUT)

    print("\n=== SUMMARY (beads per 1 mm2 square) ===")
    print("\t".join(HDR))
    for row in summary_rows:
        print("\t".join(str(row.get(h, "")) for h in HDR))
    print(f"\nWrote CSVs + xlsx + annotated images to: {OUT}")
    print(f"Collected {n_collected} annotated images into: {dest}")


if __name__ == "__main__":
    main()
